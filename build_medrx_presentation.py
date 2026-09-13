import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Set to 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    DARK_NAVY = RGBColor(15, 23, 42)     # #0F172A
    TEAL_BLUE = RGBColor(2, 132, 199)    # #0284C7
    LIGHT_BG = RGBColor(248, 250, 252)   # #F8FAFC
    TEXT_DARK = RGBColor(30, 41, 59)     # #1E293B
    SLATE_GRAY = RGBColor(100, 116, 139) # #64748B
    WHITE = RGBColor(255, 255, 255)
    CARD_BG = RGBColor(241, 245, 249)    # #F1F5F9
    ACCENT_GREEN = RGBColor(5, 150, 105) # #059669

    def set_slide_background(slide, color=LIGHT_BG):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, category_text="MEDRX PLATFORM PRESENTATION"):
        # Category Banner
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEAL_BLUE

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = DARK_NAVY

    # Avatar Paths
    avatar_dir = r"C:\Users\dell\.gemini\antigravity\brain\b6a55365-71d3-4b04-8f5c-97161587ccfe"
    img_zumair = os.path.join(avatar_dir, "zumair_ali_avatar_1789322812685.png")
    img_zahid = os.path.join(avatar_dir, "dr_zahid_avatar_1789322828158.png")
    img_mufeez = os.path.join(avatar_dir, "mufeez_ilyas_avatar_1789322845081.png")
    img_ramsha = os.path.join(avatar_dir, "ramsha_fatima_avatar_1789322861273.png")

    blank_layout = prs.slide_layouts[6]

    # =========================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, DARK_NAVY)

    # Hero Box
    hero = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.5))
    tf1 = hero.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "💊 MedRx Platform"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = TEAL_BLUE
    p1.space_after = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "Agentic AI Closed-Loop Prescription Management, Medication Safety & Audit Network"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_after = Pt(16)

    p3 = tf1.add_paragraph()
    p3.text = "A Multi-Agent Clinical Decision Support System with Dual-Credential Verification & Immutable Audit Integrity"
    p3.font.size = Pt(14)
    p3.font.color.rgb = SLATE_GRAY

    # Footer Metadata
    meta_box = s1.shapes.add_textbox(Inches(1.0), Inches(5.8), Inches(11.3), Inches(1.0))
    tf_m = meta_box.text_frame
    pm = tf_m.paragraphs[0]
    pm.text = "Team Lead: Engr. Zumair Ali  |  Riphah International University & PEL & NED UET  |  September 2026"
    pm.font.size = Pt(12)
    pm.font.color.rgb = TEAL_BLUE

    # =========================================================
    # SLIDE 2: TEAM INTRODUCTION (WITH AVATARS)
    # =========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Project Core Team & Engineering Lead Profiles", "TEAM INTRODUCTION")

    members = [
        {"name": "Engr. Zumair Ali", "role": "Team Lead & Systems Architect", "inst": "Lecturer BME", "uni": "Riphah International University", "img": img_zumair},
        {"name": "Dr. Muhammad Zahid", "role": "Clinical Domain Specialist", "inst": "Associate Professor", "uni": "Riphah International University", "img": img_zahid},
        {"name": "Engr. Mufeez Ilyas", "role": "Backend & AI Engineer", "inst": "Design Engineer", "uni": "PEL (Pak Elektron Limited)", "img": img_mufeez},
        {"name": "Ramsha Fatima", "role": "QA & Security Analyst", "inst": "BS Physics Student", "uni": "NED University (NED UET)", "img": img_ramsha},
    ]

    left_margin = Inches(0.8)
    card_width = Inches(2.7)
    card_gap = Inches(0.3)

    for i, m in enumerate(members):
        col_left = left_margin + i * (card_width + card_gap)
        col_top = Inches(1.8)

        # Card Container Shape
        shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col_left, col_top, card_width, Inches(5.0))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = CARD_BG
        shape.line.width = Pt(1.5)

        # Insert Avatar Image
        if os.path.exists(m['img']):
            s2.shapes.add_picture(m['img'], col_left + Inches(0.35), col_top + Inches(0.3), width=Inches(2.0))

        # Member Info Text
        tbox = s2.shapes.add_textbox(col_left + Inches(0.15), col_top + Inches(2.4), card_width - Inches(0.3), Inches(2.4))
        tf = tbox.text_frame
        tf.word_wrap = True

        p_name = tf.paragraphs[0]
        p_name.text = m['name']
        p_name.font.size = Pt(14)
        p_name.font.bold = True
        p_name.font.color.rgb = DARK_NAVY
        p_name.alignment = PP_ALIGN.CENTER

        p_role = tf.add_paragraph()
        p_role.text = m['role']
        p_role.font.size = Pt(10)
        p_role.font.bold = True
        p_role.font.color.rgb = TEAL_BLUE
        p_role.alignment = PP_ALIGN.CENTER
        p_role.space_after = Pt(8)

        p_inst = tf.add_paragraph()
        p_inst.text = f"• {m['inst']}"
        p_inst.font.size = Pt(9.5)
        p_inst.font.bold = True
        p_inst.font.color.rgb = TEXT_DARK
        p_inst.alignment = PP_ALIGN.CENTER

        p_uni = tf.add_paragraph()
        p_uni.text = m['uni']
        p_uni.font.size = Pt(9.0)
        p_uni.font.color.rgb = SLATE_GRAY
        p_uni.alignment = PP_ALIGN.CENTER

    # =========================================================
    # SLIDE 3: PROBLEM STATEMENT
    # =========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Critical Healthcare Failure Points Addressed by MedRx", "PROBLEM STATEMENT")

    probs = [
        {"title": "🚨 Preventable Adverse Drug Events (ADEs)", "desc": "Legacy paper & unintegrated EHR systems lack real-time allergen cross-checks. Clinicians miss penicillin/sulfa allergies or severe drug-drug interactions."},
        {"title": "🔓 Pharmacy Fraud & Over-Dispensing", "desc": "Single-factor patient lookup allows unauthorized prescription pickups. Outpatient pharmacies lack remaining-quantity tracking across partial refills."},
        {"title": "⏳ Stockout Delays & Unregulated Substitutions", "desc": "Medication stockouts cause critical treatment delays. Informal generic substitutions occur without standardized bioequivalence scoring or doctor consent."},
        {"title": "⚖️ Audit Non-Compliance & Black-Box AI Risks", "desc": "Standard hospital software uses mutable tables vulnerable to tampering. Commercial AI tools lack step-by-step clinical audit logging required by HIPAA."},
    ]

    for i, p in enumerate(probs):
        row = i // 2
        col = i % 2
        c_left = left_margin + col * Inches(5.9)
        c_top = Inches(1.8) + row * Inches(2.6)

        shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, Inches(5.6), Inches(2.3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = CARD_BG

        tb = s3.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.2), Inches(5.2), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True

        pt = tf.paragraphs[0]
        pt.text = p['title']
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = DARK_NAVY
        pt.space_after = Pt(6)

        pd = tf.add_paragraph()
        pd.text = p['desc']
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_DARK

    # =========================================================
    # SLIDE 4: OUR SOLUTION
    # =========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Closed-Loop AI Architecture & Dual-Credential Verification", "OUR SOLUTION")

    sol_cards = [
        {"num": "01", "title": "4-Agent AI Ecosystem", "body": "Prescription Verification, Medication Safety, Alternative Medicine, and Audit Anomaly Agents working in synergy with deterministic fallback."},
        {"num": "02", "title": "Dual-Credential Security", "body": "Unlocking prescriptions requires Prescription ID + 4-digit PIN. Public QR verification discloses ZERO Patient Health Information (PHI)."},
        {"num": "03", "title": "Human-in-the-Loop Approval", "body": "Pharmacist out-of-stock alternatives are scored (0-100%) and routed to the prescribing doctor. No substitution occurs without explicit sign-off."},
        {"num": "04", "title": "Immutable Audit Trail", "body": "Append-only log architecture redacting secrets metadata (`sanitize_audit_metadata`). Complete compliance with GxP & HIPAA audit standards."},
    ]

    for i, sc in enumerate(sol_cards):
        c_left = left_margin + i * Inches(2.95)
        c_top = Inches(1.8)

        shape = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, Inches(2.75), Inches(5.0))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = TEAL_BLUE
        shape.line.width = Pt(1)

        tb = s4.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.3), Inches(2.35), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True

        pnum = tf.paragraphs[0]
        pnum.text = sc['num']
        pnum.font.size = Pt(28)
        pnum.font.bold = True
        pnum.font.color.rgb = TEAL_BLUE
        pnum.space_after = Pt(8)

        ptit = tf.add_paragraph()
        ptit.text = sc['title']
        ptit.font.size = Pt(14)
        ptit.font.bold = True
        ptit.font.color.rgb = DARK_NAVY
        ptit.space_after = Pt(10)

        pbody = tf.add_paragraph()
        pbody.text = sc['body']
        pbody.font.size = Pt(10)
        pbody.font.color.rgb = TEXT_DARK

    # =========================================================
    # SLIDE 5: MULTI-AGENT AI ECOSYSTEM FLOW DIAGRAM
    # =========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Multi-Agent AI Decision Support & Safety Framework", "AI AGENT ARCHITECTURE")

    agents_list = [
        {"title": "Agent 1: Verification Agent", "desc": "Evaluates dosage units, strength values, and administration route appropriateness upon order submission."},
        {"title": "Agent 2: Medication Safety Agent", "desc": "Cross-references patient allergen arrays (Penicillin, Sulfa) & chronic diseases against active ingredients."},
        {"title": "Agent 3: Alternative Medicine Agent", "desc": "Scores bioequivalence (0-100%) for stockouts and structures substitution rationale for physician approval."},
        {"title": "Agent 4: Anomaly Detection Agent", "desc": "Scans append-only logs for failed PIN spikes, abnormal dispensing velocity, and unauthorized access attempts."},
    ]

    for i, ag in enumerate(agents_list):
        c_left = left_margin + i * Inches(2.95)
        c_top = Inches(2.0)

        shape = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, Inches(2.75), Inches(4.6))
        shape.fill.solid()
        shape.fill.fore_color.rgb = DARK_NAVY

        tb = s5.shapes.add_textbox(c_left + Inches(0.15), c_top + Inches(0.4), Inches(2.45), Inches(4.0))
        tf = tb.text_frame
        tf.word_wrap = True

        pt = tf.paragraphs[0]
        pt.text = ag['title']
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEAL_BLUE
        pt.space_after = Pt(10)

        pd = tf.add_paragraph()
        pd.text = ag['desc']
        pd.font.size = Pt(10)
        pd.font.color.rgb = LIGHT_BG

    # =========================================================
    # SLIDE 6: CLINICAL WORKFLOW
    # =========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Prescribe → Verify → Dispense → Approve → Audit Workflow", "CLINICAL WORKFLOW")

    steps = [
        {"step": "STEP 1", "title": "Doctor Prescribes", "desc": "Doctor completes Encounter Studio (vitals, ICD-10) and issues prescription with 4-digit PIN."},
        {"step": "STEP 2", "title": "AI Safety Scan", "desc": "Verification & Safety Agents evaluate dosing & allergen overlap in real time."},
        {"step": "STEP 3", "title": "Pharmacy Dispensing", "desc": "Pharmacist unlocks Rx via ID + PIN. Dispensing Guard locks remaining quantity."},
        {"step": "STEP 4", "title": "Doctor Approval", "desc": "Stockouts trigger Alternative Agent. Doctor approves substitution in Inbox (v1 -> v2)."},
    ]

    for i, stp in enumerate(steps):
        c_left = left_margin + i * Inches(2.95)
        c_top = Inches(2.0)

        shape = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, Inches(2.75), Inches(4.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = CARD_BG

        tb = s6.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.3), Inches(2.35), Inches(3.9))
        tf = tb.text_frame
        tf.word_wrap = True

        ps = tf.paragraphs[0]
        ps.text = stp['step']
        ps.font.size = Pt(12)
        ps.font.bold = True
        ps.font.color.rgb = TEAL_BLUE
        ps.space_after = Pt(4)

        pt = tf.add_paragraph()
        pt.text = stp['title']
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = DARK_NAVY
        pt.space_after = Pt(8)

        pd = tf.add_paragraph()
        pd.text = stp['desc']
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_DARK

    # =========================================================
    # SLIDE 7: KEY INNOVATIONS
    # =========================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Architectural Integrity & Advanced Engineering Features", "KEY INNOVATIONS")

    innovs = [
        {"title": "Single-Increment Versioning", "desc": "Approved doctor substitutions increment prescription version (v1 -> v2) while leaving original records intact for audit history."},
        {"title": "Fallback Doctor Routing", "desc": "If prescribing doctor is inactive/suspended, requests route automatically to hospital on-call doctors or administrative escalation."},
        {"title": "Duplicate Submission Protection", "desc": "5-minute short-window duplicate warning banner and unique submission tokens (`submission_token`) prevent double-issuance."},
        {"title": "SQLite WAL Mode Concurrency", "desc": "Configured Write-Ahead Logging (`PRAGMA journal_mode=WAL; PRAGMA busy_timeout=5000;`) for high-concurrency database safety."},
    ]

    for i, inv in enumerate(innovs):
        row = i // 2
        col = i % 2
        c_left = left_margin + col * Inches(5.9)
        c_top = Inches(1.8) + row * Inches(2.6)

        shape = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, Inches(5.6), Inches(2.3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = CARD_BG

        tb = s7.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.2), Inches(5.2), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True

        pt = tf.paragraphs[0]
        pt.text = inv['title']
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = DARK_NAVY
        pt.space_after = Pt(6)

        pd = tf.add_paragraph()
        pd.text = inv['desc']
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_DARK

    # =========================================================
    # SLIDE 8: SECURITY & PRIVACY
    # =========================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Cryptographic Protections & HIPAA-Aligned Privacy", "SECURITY & PRIVACY")

    secs = [
        {"title": "🔒 Zero-PHI QR Verification", "desc": "QR code contains ONLY an opaque URL string (`https://medrx.health/verify?token=...`). Zero patient names, DOBs, CNICs, or diagnoses are embedded in the QR."},
        {"title": "🔑 Constant-Time Comparison", "desc": "Uses `hmac.compare_digest` for security PIN and password verification, rendering timing attack vulnerabilities impossible."},
        {"title": "🛡️ SHA-256 & PBKDF2 Hashing", "desc": "Verification tokens and 4-digit PINs are stored as SHA-256 / PBKDF2 cryptographic hashes. Zero plaintext PINs stored in the database."},
        {"title": "📜 Secret Redaction in Audit Logs", "desc": "Audit logging engine automatically sanitizes passwords, PINs, and API keys (`sanitize_audit_metadata`) prior to log creation."},
    ]

    for i, sc in enumerate(secs):
        row = i // 2
        col = i % 2
        c_left = left_margin + col * Inches(5.9)
        c_top = Inches(1.8) + row * Inches(2.6)

        shape = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, Inches(5.6), Inches(2.3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = CARD_BG

        tb = s8.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.2), Inches(5.2), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True

        pt = tf.paragraphs[0]
        pt.text = sc['title']
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = DARK_NAVY
        pt.space_after = Pt(6)

        pd = tf.add_paragraph()
        pd.text = sc['desc']
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_DARK

    # =========================================================
    # SLIDE 9: EVIDENCES & QA MATRIX
    # =========================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "100% Automated Test Suite Verification Results", "EVIDENCES & QA METRICS")

    # Big Banner Result
    b_shape = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_margin, Inches(1.8), Inches(11.733), Inches(1.2))
    b_shape.fill.solid()
    b_shape.fill.fore_color.rgb = ACCENT_GREEN

    tb_b = s9.shapes.add_textbox(left_margin + Inches(0.2), Inches(1.9), Inches(11.333), Inches(1.0))
    tf_b = tb_b.text_frame
    pb = tf_b.paragraphs[0]
    pb.text = "✅ 98 / 98 AUTOMATED TESTS PASSED (0 FAILURES, 0 ERRORS)"
    pb.font.size = Pt(18)
    pb.font.bold = True
    pb.font.color.rgb = WHITE
    pb.alignment = PP_ALIGN.CENTER

    pb_sub = tf_b.add_paragraph()
    pb_sub.text = "Full system regression suite executed cleanly across all 5 test modules (`run_all_tests.py`)"
    pb_sub.font.size = Pt(12)
    pb_sub.font.color.rgb = WHITE
    pb_sub.alignment = PP_ALIGN.CENTER

    # Test Table Summary
    qa_list = [
        {"mod": "test_app.py", "scope": "Core DB, RBAC, Encounters, Prescriptions & PDF Generation", "res": "40 / 40 Passed"},
        {"mod": "test_phase7.py", "scope": "Substitution Approval Engine & Fallback Doctor Routing", "res": "7 / 7 Passed"},
        {"mod": "test_phase8.py", "scope": "System Audit Logging & Anomaly Detection Rules A-F", "res": "15 / 15 Passed"},
        {"mod": "test_phase9.py", "scope": "QR Verification, Security Hardening & Zero-PHI Audit", "res": "16 / 16 Passed"},
        {"mod": "test_phase10.py", "scope": "Version Idempotency, Concurrency & UI Credentials Handover", "res": "20 / 20 Passed"},
    ]

    for i, q in enumerate(qa_list):
        c_top = Inches(3.2) + i * Inches(0.75)
        shape = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_margin, c_top, Inches(11.733), Inches(0.65))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE if i % 2 == 0 else CARD_BG
        shape.line.color.rgb = CARD_BG

        tb = s9.shapes.add_textbox(left_margin + Inches(0.2), c_top + Inches(0.05), Inches(11.333), Inches(0.55))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = f"• {q['mod']}  |  {q['scope']}  --->  {q['res']}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = DARK_NAVY

    # =========================================================
    # SLIDE 10: CONCLUSION & SUMMARY
    # =========================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, DARK_NAVY)

    tb_end = s10.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.5))
    tf_e = tb_end.text_frame
    tf_e.word_wrap = True

    pe1 = tf_e.paragraphs[0]
    pe1.text = "Thank You!"
    pe1.font.size = Pt(40)
    pe1.font.bold = True
    pe1.font.color.rgb = TEAL_BLUE
    pe1.space_after = Pt(14)

    pe2 = tf_e.add_paragraph()
    pe2.text = "MedRx provides a zero-compromise platform for prescription safety, multi-agent AI verification, and regulatory compliance."
    pe2.font.size = Pt(18)
    pe2.font.bold = True
    pe2.font.color.rgb = WHITE
    pe2.space_after = Pt(24)

    pe3 = tf_e.add_paragraph()
    pe3.text = "Live Streamlit Cloud URL: https://medrx.streamlit.app\nGitHub Source Repository: https://github.com/YOUR_USERNAME/medrx"
    pe3.font.size = Pt(13)
    pe3.font.color.rgb = SLATE_GRAY

    # Save Presentation
    output_filename = "MedRx_Project_Presentation.pptx"
    prs.save(output_filename)
    print(f"Successfully generated PowerPoint presentation: {output_filename}")

if __name__ == "__main__":
    create_presentation()
